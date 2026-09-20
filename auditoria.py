#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Auditoría de calidad del banco de Solaris Merit Neiva.

Mide solo el banco principal. Las preguntas marcadas con tipo_v2 = "repaso" quedan fuera
de los indicadores: son el cuestionario de nivel fácil, de opciones cortas y directas.

Detecta las pistas que permiten acertar sin saber la norma. Se corre antes de publicar
cualquier tanda nueva de preguntas:

    python3 auditoria.py banco-preguntas.xlsx

Revisa cinco cosas:
  1. Longitud: que la opción correcta no sea sistemáticamente la más larga.
  2. Absolutos: "nunca", "siempre", "ninguna" en los distractores, que se descartan solos.
  3. Lenguaje calificado: que los matices ("conforme a", "por escrito", "dejando
     constancia") estén repartidos y no solo en la correcta.
  4. Solapamiento léxico: que la correcta no repita las palabras del enunciado.
  5. Balance de la clave a/b/c y de la dificultad.
"""
import sys, re, statistics as st
from collections import Counter
from openpyxl import load_workbook

RUTA = sys.argv[1] if len(sys.argv) > 1 else "banco-preguntas.xlsx"
MATICES = r"\b(conforme|según|sin perjuicio|salvo|documenta\w*|por escrito|registra\w*|informa\w*|verifica\w*|constancia|dentro de los términos|en el marco)\b"
ABSOLUTOS = r"\b(nunca|siempre|ninguna|ningún|jamás|en ningún caso|todos los casos)\b"

wb = load_workbook(RUTA, data_only=True)
ws = wb["Funcionales"]
cols = [c.value for c in ws[1]]
TODAS = [dict(zip(cols, r)) for r in ws.iter_rows(min_row=2, values_only=True)
         if r[0] and str(r[cols.index("activa")]).upper() != "NO"]
REPASO = [q for q in TODAS if str(q.get("tipo_v2") or "").strip() == "repaso"]
F = [q for q in TODAS if q not in REPASO]

def clave(q): return str(q["respuesta"]).strip().lower()
def otras(q): return [k for k in "abc" if k != clave(q)]
def palabras(t): return set(w.lower() for w in re.findall(r"\w{6,}", str(t)))

alertas = []
largo_c, largo_o, gaps, mas_larga = [], [], [], 0
abs_dist, matiz_c, matiz_o, matiz_solo, solape = 0, 0, 0, 0, 0

for q in F:
    k = clave(q)
    lc = len(str(q[k])); lo = [len(str(q[x])) for x in otras(q)]
    largo_c.append(lc); largo_o += lo
    g = lc - max(lo); gaps.append(g)
    if g > 0: mas_larga += 1
    if g > 12: alertas.append((q["id"], "la correcta es %d caracteres más larga" % g))
    for x in otras(q):
        if re.search(ABSOLUTOS, str(q[x]), re.I):
            abs_dist += 1; alertas.append((q["id"], "absoluto en la opción %s" % x))
        if re.search(MATICES, str(q[x]), re.I): matiz_o += 1
    if re.search(MATICES, str(q[k]), re.I):
        matiz_c += 1
        if not any(re.search(MATICES, str(q[x]), re.I) for x in otras(q)):
            matiz_solo += 1
            alertas.append((q["id"], "solo la correcta usa lenguaje calificado"))
    pe = palabras(q["enunciado"])
    if len(pe & palabras(q[k])) > max(len(pe & palabras(q[x])) for x in otras(q)):
        solape += 1

n = len(F)
print(f"BANCO PRINCIPAL: {n} preguntas · {len(set(q['caso_id'] for q in F))} casos · {len(set(q['eje'] for q in F))} ejes")
print(f"REPASO (nivel fácil, banco aparte): {len(REPASO)} preguntas · {len(set(q['caso_id'] for q in REPASO))} casos\n")
print("1. LONGITUD")
print(f"   correcta {st.mean(largo_c):.0f} car · incorrectas {st.mean(largo_o):.0f} car · ventaja media {st.mean(gaps):+.0f}")
print(f"   la correcta es la más larga en {mas_larga/n*100:.0f} %   (meta: menos de 50 %)")
print("2. ABSOLUTOS EN DISTRACTORES")
print(f"   {abs_dist} de {n*2} opciones incorrectas   (meta: menos del 3 %)")
print("3. LENGUAJE CALIFICADO")
print(f"   matiz en la correcta {matiz_c/n*100:.0f} % · en las incorrectas {matiz_o/(n*2)*100:.0f} %")
print(f"   solo en la correcta: {matiz_solo/n*100:.0f} %   (meta: 0 %)")
print("4. SOLAPAMIENTO CON EL ENUNCIADO")
print(f"   la correcta repite más palabras del enunciado en {solape/n*100:.0f} %   (meta: menos de 40 %)")
print("5. BALANCE")
print("   clave:", dict(Counter(clave(q) for q in F)))
print("   dificultad:", dict(Counter(str(q.get("dificultad") or "sin marcar") for q in F)))

if "Comportamentales" in wb.sheetnames:
    wc = wb["Comportamentales"]; cc = [c.value for c in wc[1]]
    B = [dict(zip(cc, r)) for r in wc.iter_rows(min_row=2, values_only=True)
         if r[0] and str(r[cc.index("activa")]).upper() != "NO"]
    g = []
    for q in B:
        pts = {k: int(q["puntos_" + k]) for k in "abc"}
        m = max(pts, key=pts.get)
        g.append(len(str(q[m])) - max(len(str(q[k])) for k in "abc" if k != m))
    print(f"\nCOMPORTAMENTALES: {len(B)} situaciones · ventaja de la de 3 puntos {st.mean(g):+.0f} car")

print(f"\nALERTAS: {len(alertas)}")
for i, (rid, txt) in enumerate(alertas[:25]):
    print(f"   {rid}: {txt}")
if len(alertas) > 25:
    print(f"   ... y {len(alertas)-25} más")
