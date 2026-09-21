#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Auditoría adversarial del banco.

La idea: un resolvedor que NO lee el caso ni el enunciado y solo mira las tres opciones,
aplicando las heurísticas que usa un aspirante entrenado en exámenes:
  H1  la más larga
  H2  la que trae más lenguaje calificado ("conforme a", "por escrito", "dejando constancia")
  H3  descartar las que traen absolutos y, entre las que quedan, la más larga
  H4  la que menciona informar, reportar, documentar o avisar al jefe
  H5  la que encadena dos acciones con "y"
  H6  combinación ponderada de todas
Si alguna supera el 33 % de azar por un margen amplio, el banco se puede aprobar sin saber.
Se reporta el resultado global, por nivel de dificultad, por nivel del empleo y por eje.
"""
import sys, re, random
from collections import defaultdict, Counter
from openpyxl import load_workbook

RUTA = sys.argv[1] if len(sys.argv) > 1 else "banco-preguntas.xlsx"
random.seed(11)
MAT = r"\b(conforme|según|sin perjuicio|salvo|documenta\w*|por escrito|registra\w*|informa\w*|verifica\w*|constancia|dentro de los términos|en el marco)\b"
ABS = r"\b(nunca|siempre|ninguna|ningún|jamás|en ningún caso|todos los casos|únicamente)\b"
ESC = r"\b(informa\w*|report\w*|jefe|superior|comit[eé]|documenta\w*|constancia)\b"

wb = load_workbook(RUTA, data_only=True)
ws = wb["Funcionales"]; cols = [c.value for c in ws[1]]
TODAS = [dict(zip(cols, r)) for r in ws.iter_rows(min_row=2, values_only=True)
         if r[0] and str(r[cols.index("activa")]).upper() != "NO"]
P = [q for q in TODAS if str(q.get("tipo_v2") or "").strip() != "repaso"]
R = [q for q in TODAS if str(q.get("tipo_v2") or "").strip() == "repaso"]

def clave(q): return str(q["respuesta"]).strip().lower()
def ops(q):
    """Baraja las opciones como lo hace la app y devuelve (opciones, letra correcta),
    para que el resolvedor no se beneficie del orden en que están en el Excel."""
    textos = [(k, str(q[k])) for k in "abc"]
    random.shuffle(textos)
    m = {}; correcta = None
    for nueva, (vieja, txt) in zip("abc", textos):
        m[nueva] = txt
        if vieja == clave(q): correcta = nueva
    return m, correcta

def desempata(cand):
    """Entre opciones empatadas elige al azar, como haría quien aplica la heurística."""
    return random.choice(cand) if cand else None
def h_larga(o):
    m = max(len(v) for v in o.values())
    return desempata([k for k, v in o.items() if len(v) == m])
def h_matiz(o):
    c = {k: len(re.findall(MAT, v, re.I)) for k, v in o.items()}
    m = max(c.values())
    return desempata([k for k in c if c[k] == m]) if m > 0 and list(c.values()).count(m) < 3 else None
def h_sin_abs(o):
    v = {k: t for k, t in o.items() if not re.search(ABS, t, re.I)}
    if not v: return None
    m = max(len(t) for t in v.values())
    return desempata([k for k, t in v.items() if len(t) == m])
def h_escala(o):
    c = {k: len(re.findall(ESC, v, re.I)) for k, v in o.items()}
    m = max(c.values())
    return desempata([k for k in c if c[k] == m]) if m > 0 and list(c.values()).count(m) < 3 else None
def h_encadena(o):
    c = {k: v.count(" y ") for k, v in o.items()}
    m = max(c.values())
    return desempata([k for k in c if c[k] == m]) if m > 0 and list(c.values()).count(m) < 3 else None
def h_combo(o):
    p = {}
    for k, v in o.items():
        s = len(v)/100
        s += 1.2*len(re.findall(MAT, v, re.I))
        s += 0.8*len(re.findall(ESC, v, re.I))
        s += 0.6*v.count(" y ")
        if re.search(ABS, v, re.I): s -= 2.5
        p[k] = s
    m = max(p.values())
    return desempata([k for k in p if p[k] == m])

def _finales(lote):
    """Detecta terminaciones repetidas en muchas opciones (coletillas de plantilla)."""
    from collections import Counter
    c = Counter()
    for q in lote:
        for x in "abc":
            w = str(q[x]).rstrip(".").split()
            for n in (4, 5, 6):
                if len(w) > n + 2: c[" ".join(w[-n:]).lower()] += 1
    return {f for f, n in c.items() if n >= 6}
FINALES = set()
def h_coletilla(o):
    """Descarta las opciones que terminan en una coletilla repetida y escoge entre las demás."""
    marcadas = [k for k, v in o.items() if any(v.rstrip(".").lower().endswith(f) for f in FINALES)]
    if not marcadas or len(marcadas) == 3: return None
    return desempata([k for k in o if k not in marcadas])

def h_punto(o):
    """Escoge la única opción que termina en punto, si la hay."""
    c = [k for k, v in o.items() if v.rstrip().endswith(".")]
    return c[0] if len(c) == 1 else None

H = [("H1 la más larga", h_larga), ("H2 más lenguaje calificado", h_matiz),
     ("H3 sin absolutos y más larga", h_sin_abs), ("H4 la que informa o documenta", h_escala),
     ("H5 la que encadena dos acciones", h_encadena), ("H6 combinación de todas", h_combo),
     ("H7 descarta la de coletilla", h_coletilla), ("H8 la única con punto final", h_punto)]

def evalua(lote, nombre):
    if not lote: return
    global FINALES
    FINALES = _finales(lote)
    print(f"\n{nombre}  ({len(lote)} preguntas · azar 33 %)")
    peor = 0
    for etq, f in H:
        ac = tot = 0
        for q in lote:
            o, cor = ops(q)
            p = f(o)
            if p is None: continue
            tot += 1
            if p == cor: ac += 1
        if tot:
            pct = ac/tot*100
            peor = max(peor, pct)
            flag = "  ← DELATA" if pct >= 50 else ("  ← revisar" if pct >= 42 else "")
            print(f"   {etq:<34} {pct:5.1f} %   (aplica en {tot/len(lote)*100:.0f} % de los ítems){flag}")
    print(f"   PEOR HEURÍSTICA: {peor:.1f} %")
    return peor

print("="*74)
print("AUDITORÍA ADVERSARIAL · SOLARIS MERIT NEIVA")
print("="*74)
print(f"Banco principal: {len(P)} preguntas · {len(set(q['caso_id'] for q in P))} casos")
print(f"Repaso:          {len(R)} preguntas · {len(set(q['caso_id'] for q in R))} casos")

evalua(P, "BANCO PRINCIPAL")
evalua(R, "BANCO DE REPASO (nivel fácil, se espera que sea adivinable)")

print("\n" + "-"*74)
print("BALANCE DE LA CLAVE EN EL EXCEL (sin barajar, para quien estudie desde el archivo)")
for nom, lote in (("principal", P), ("repaso", R)):
    c = Counter(clave(q) for q in lote)
    peor = max(c.values())/len(lote)*100 if lote else 0
    print(f"   {nom:<10} {dict(c)} · marcar siempre la letra más frecuente: {peor:.1f} %" + ("  ← DELATA" if peor >= 45 else ""))
print("\n" + "-"*74)
print("POR NIVEL DE DIFICULTAD")
for d in ["Basica", "Media", "Alta"]:
    evalua([q for q in P if str(q.get("dificultad") or "").strip() == d], f"dificultad {d}")
sinmarcar = [q for q in P if not str(q.get("dificultad") or "").strip()]
if sinmarcar: evalua(sinmarcar, "sin dificultad marcada")

print("\n" + "-"*74)
print("POR NIVEL DEL EMPLEO")
for n in ["Profesional", "Tecnico"]:
    evalua([q for q in P if str(q.get("nivel") or "").strip() == n], f"nivel {n}")

print("\n" + "-"*74)
print("PEOR HEURÍSTICA POR EJE (principal)")
por = defaultdict(list)
for q in P: por[q["eje"]].append(q)
fila = []
for e, lot in por.items():
    mejor = 0
    for etq, f in H:
        ac = tot = 0
        for q in lot:
            o, cor = ops(q)
            p = f(o)
            if p is None: continue
            tot += 1
            if p == cor: ac += 1
        if tot: mejor = max(mejor, ac/tot*100)
    fila.append((mejor, len(lot), e))
for m, n, e in sorted(fila, reverse=True):
    marca = "  ← DELATA" if m >= 50 else ""
    print(f"   {m:5.1f} %  ({n:>3} preg.)  {e}{marca}")

# comportamentales
if "Comportamentales" in wb.sheetnames:
    wc = wb["Comportamentales"]; cc = [c.value for c in wc[1]]
    B = [dict(zip(cc, r)) for r in wc.iter_rows(min_row=2, values_only=True)
         if r[0] and str(r[cc.index("activa")]).upper() != "NO"]
    print("\n" + "-"*74)
    print(f"COMPORTAMENTALES ({len(B)} situaciones · azar 33 %)")
    for etq, f in H:
        ac = tot = 0
        for q in B:
            pts = {k: int(q["puntos_" + k]) for k in "abc"}
            mejor_orig = max(pts, key=pts.get)
            textos = [(k, str(q[k])) for k in "abc"]; random.shuffle(textos)
            o = {}; bueno = None
            for nueva, (vieja, txt) in zip("abc", textos):
                o[nueva] = txt
                if vieja == mejor_orig: bueno = nueva
            p = f(o)
            if p is None: continue
            tot += 1
            if p == bueno: ac += 1
        if tot:
            pct = ac/tot*100
            flag = "  ← DELATA" if pct >= 50 else ("  ← revisar" if pct >= 42 else "")
            print(f"   {etq:<34} {pct:5.1f} %   (aplica en {tot/len(B)*100:.0f} %){flag}")

print("\n" + "="*74)
